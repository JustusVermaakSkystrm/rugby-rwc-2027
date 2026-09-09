#!/usr/bin/env python3
"""Rewrite slide 3 of the Codi deck in plainer language.

Works from Sandra's returned file so her own edits to slides 1, 2 and 4 are
preserved. Every text box on slide 3 is a single paragraph with a single run,
so swapping run.text keeps all formatting (font, size, colour, spacing).
"""
import sys
from pptx import Presentation

REPLACEMENTS = {
    # --- headings ---
    "Three hypotheses I would arrive with":
        "Three things I'd check first",

    "Each is a question with a data test attached — not a conclusion.":
        "These are questions, not answers. Each one has something I can go and check.",

    # --- card headings ---
    "Delivery mix": "Who does the work",
    "Regulation":   "The new repair rules",
    "Price":        "Costs going up",

    # --- card evidence ---
    "Pobl expanded its in-house trades team by c.50% — 34 roles — in 2024, "
    "alongside an 11-lot framework for responsive repairs, voids and out-of-hours.":
        "Pobl grew its in-house trades team by about half (34 roles) in 2024. "
        "There is also a contractor framework of 11 lots covering repairs, voids "
        "and out-of-hours.",

    "WHQS hazard response took effect 1 April 2026: investigate and remedy within "
    "24 hours each where harm is imminent; 10 then 5 working days otherwise.":
        "From 1 April 2026, WHQS sets response times. Urgent hazards: 24 hours to "
        "investigate, 24 to fix. Everything else: 10 working days, then 5.",

    "Sector surveys put expected responsive repairs cost increases at c.23% this "
    "year, against c.10% in the regulator's global accounts.":
        "Repairs costs across the sector are expected to rise by about 23% this "
        "year. The regulator's published figures suggested nearer 10%.",

    # --- label ---
    "THE TEST": "WHAT I'D CHECK",

    # --- the checks ---
    "Has external spend stepped down in proportion to the capacity we insourced, "
    "or are we carrying both?":
        "Has contractor spend come down now we do more in-house, or are we paying "
        "for both?",

    "What has that done to job volume, out-of-hours and premium call-off, and to "
    "how densely we can schedule?":
        "How much extra out-of-hours and emergency work has this created, and is "
        "it harder to plan jobs?",

    "On like-for-like jobs, how much of our movement is rate and materials rather "
    "than volume?":
        "For the same type of job, how much of our increase is price rather than "
        "more jobs?",

    # --- sources ---
    "Sources: Codi/Pobl published announcements and tender notices; Welsh "
    "Government WHQS hazard-response statement; UK sector cost surveys.":
        "Sources: Pobl announcements and tender notices; Welsh Government WHQS "
        "statement; UK-wide sector cost surveys.",
}

NOTES = """PACE 2:05  ·  Running total 4:10.

I wouldn't turn up with no idea what was going on. I'd come with three things I want to check, and the data I'd need to prove or rule out each one.

First, who does the work. Pobl grew its in-house trades team by about half in 2024, thirty-four roles, and there's also a contractor framework across eleven lots covering repairs, voids and out-of-hours. That's a sensible set-up, but it's exactly where you can end up paying twice. So I'd check whether contractor spend has come down in line with the work we brought in-house, comparing it month by month against the jobs our own teams complete.

Second, the new repair rules. Since the first of April this year, WHQS sets response times. If a hazard is urgent, it's twenty-four hours to investigate and another twenty-four to put it right. Otherwise it's ten working days, then five. Tighter deadlines mean more out-of-hours and more emergency call-outs, and they make it harder to group jobs together sensibly, so the cost per job goes up even when the number of jobs doesn't. And because response times have to be published, this isn't somewhere to look for savings.

Third, costs going up. Across the sector, repairs costs are expected to rise by about twenty-three per cent this year, against nearer ten in the regulator's published figures. That's UK-wide rather than Welsh, so I'd use it as a reason to separate price from volume, not as a number to lean on.

If none of the three is right, I've still narrowed it down quickly. I'd rather be put right in week one than be confident and wrong in month six.

[Source] Pobl in-house trades announcement (2024); Pobl responsive repairs and voids framework notice, 11 lots (Dec 2023); Welsh Government WHQS responding-to-hazards statement (Dec 2025), in force 1 April 2026; UK repairs cost survey commentary. Context, not claims about Codi's position."""


def main(src, dst):
    prs = Presentation(src)
    slide = prs.slides[2]

    hits, misses = 0, dict(REPLACEMENTS)
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            if not para.runs:
                continue
            current = "".join(r.text for r in para.runs)
            if current in REPLACEMENTS:
                para.runs[0].text = REPLACEMENTS[current]
                for extra in para.runs[1:]:
                    extra.text = ""
                hits += 1
                misses.pop(current, None)

    slide.notes_slide.notes_text_frame.text = NOTES
    prs.save(dst)

    print(f"replaced {hits} text blocks on slide 3")
    if misses:
        print("NOT FOUND (check wording):")
        for k in misses:
            print("  -", k[:70])
    words = len(NOTES.split("\n", 1)[1].split("[Source]")[0].split())
    print(f"slide 3 notes: {words} words -> {words/130*60:.0f}s at 130wpm (target 125s)")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
