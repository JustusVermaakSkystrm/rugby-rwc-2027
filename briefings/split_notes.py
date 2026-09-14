#!/usr/bin/env python3
"""Separate the speaker notes from the deck.

Writes:
  * a notes-free .pptx, with the notes parts removed from the package entirely
    (not just blanked), and the employer sensitivity label stripped
  * notes.json, the extracted text, for building the Word file

Authorship and editing history — author, created, modified, revision and
TotalTime — are deliberately left untouched.

Usage: split_notes.py <in.pptx> <out.pptx> <notes.json>
"""
import json
import re
import shutil
import sys
import zipfile

from pptx import Presentation

REMOVED = []


def extract(src):
    prs = Presentation(src)
    out = []
    for n, slide in enumerate(prs.slides, 1):
        title = ""
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                top = sh.top / 914400
                if 0.6 <= top < 1.4:
                    title = sh.text_frame.text.strip()
                    break
        if not title:
            title = "Title slide"
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        cue, body = "", notes
        m = re.match(r"\[(.*?)\]\s*\n+(.*)", notes, re.S)
        if m:
            cue, body = m.group(1).strip(), m.group(2).strip()
        out.append({"slide": n, "title": title, "cue": cue,
                    "paragraphs": [p.strip() for p in body.split("\n\n") if p.strip()]})
    return out


def strip_package(src, dst):
    """Drop every notesSlide part, its relationships, and the sensitivity label."""
    drop_prefixes = ["ppt/notesSlides/", "ppt/notesMasters/", "docMetadata/"]
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        names = set(zin.namelist())

        # the notes master owns its own theme part; removing the master
        # without it leaves an unreferenced theme, which PowerPoint reports
        # as a corrupt file
        extra_drop = set()
        for rels in [n for n in names if re.match(r"ppt/notesMasters/_rels/.*\.rels$", n)]:
            for target in re.findall(r'Target="([^"]+)"', zin.read(rels).decode("utf8")):
                if "theme" in target:
                    extra_drop.add("ppt/" + target.replace("../", ""))
        for item in zin.infolist():
            name = item.filename

            if name.startswith(tuple(drop_prefixes)) or name in extra_drop:
                REMOVED.append(name)
                continue
            # the _rels files that belong to dropped parts
            if name.startswith("ppt/notesSlides/_rels") or name.startswith("ppt/notesMasters/_rels"):
                REMOVED.append(name)
                continue

            data = zin.read(name)

            if name == "[Content_Types].xml":
                text = data.decode("utf8")
                text = re.sub(r'<Override[^>]*PartName="/(ppt/notesSlides|ppt/notesMasters|docMetadata)[^>]*/>',
                              "", text)
                for orphan in extra_drop:
                    text = re.sub(r'<Override[^>]*PartName="/%s"[^>]*/>' % re.escape(orphan),
                                  "", text)
                data = text.encode("utf8")

            if name == "_rels/.rels":
                text = data.decode("utf8")
                text = re.sub(r'<Relationship[^>]*classificationlabels[^>]*/>', "", text)
                data = text.encode("utf8")

            # slide -> notesSlide references, and presentation -> notesMaster
            if re.match(r"ppt/slides/_rels/slide\d+\.xml\.rels$", name) or \
               name == "ppt/_rels/presentation.xml.rels":
                text = data.decode("utf8")
                text = re.sub(r'<Relationship[^>]*(notesSlide|notesMaster)[^>]*/>', "", text)
                data = text.encode("utf8")

            if name == "ppt/presentation.xml":
                text = data.decode("utf8")
                text = re.sub(r"<p:notesMasterIdLst>.*?</p:notesMasterIdLst>", "", text, flags=re.S)
                text = re.sub(r"<p:notesMasterIdLst\s*/>", "", text)
                data = text.encode("utf8")

            zout.writestr(item, data)


def main(src, dst, notes_json):
    notes = extract(src)
    with open(notes_json, "w") as f:
        json.dump(notes, f, indent=1)

    tmp = dst + ".tmp"
    strip_package(src, tmp)
    shutil.move(tmp, dst)

    # confirm python-pptx still opens it and the notes really are gone
    prs = Presentation(dst)
    left = sum(1 for s in prs.slides
               if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    words = sum(len(" ".join(n["paragraphs"]).split()) for n in notes)

    print(f"extracted {len(notes)} sets of notes ({words} words) -> {notes_json}")
    print(f"removed {len(REMOVED)} parts from the package:")
    for r in sorted(REMOVED):
        print(f"    {r}")
    print(f"\nnotes remaining in {dst}: {left}  (expected 0)")
    print(f"slides: {len(prs.slides)}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
